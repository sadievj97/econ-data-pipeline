"""Files downloaded by hand, kept in data/manual/ with a note on where they came from.

A manual file is a CSV, or an .xlsx sheet (filters: {sheet: ...}), with one row
per country and year: columns iso3 and year, one column per series (the spec's
code names the column), and optionally a column called note. Each file needs
a note beside it named <file>.source.md, saying where and when it was
downloaded. The file is copied into data/raw/ like any download; later builds
use that copy, and stop if the file in data/manual/ has changed since.
Units cannot be read from such a file, so the spec's units are used and the
fetch says so.
"""

from __future__ import annotations

import csv
import io
import math
import re
import time
from collections.abc import Callable, Iterable
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

from ..cache import RawCache, Snapshot, sha256_bytes
from ..spec import SourceRef
from ..units import UnitInfo
from .base import FetchedSeries, IndicatorInfo, SourceError, make_frame

UNCHECKED = "units are as declared in the spec: a hand-downloaded file does not state them, so check them against the source"


def _safe(part: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", part)


class Manual:
    name = "manual"

    def __init__(self, cache: RawCache, client: httpx.Client | None = None, attempts: int = 4,
                 sleep: Callable[[float], None] = time.sleep, manual_dir: Path | None = None, **_: Any):
        self.cache = cache
        self.manual_dir = manual_dir

    def catalog(self, refresh: bool = False) -> dict[str, IndicatorInfo]:
        return {}

    def _snapshot(self, ref: SourceRef, refresh: bool) -> Snapshot:
        if self.manual_dir is None:
            raise SourceError("no data/manual folder is configured")
        path = self.manual_dir / Path(ref.file or "")
        if not path.is_file():
            raise SourceError(f"data/manual/{ref.file} does not exist")
        note = path.with_name(path.name + ".source.md")
        if not note.is_file():
            raise SourceError(f"data/manual/{ref.file} has no note saying where it came from: add {note.name} next to it "
                              "(see data/manual/README.md)")
        content = path.read_bytes()
        parts = [_safe(p) for p in Path(ref.file or "").parts]
        key = "manual/" + "/".join(parts)
        latest = self.cache.latest(key)
        if latest is not None and not refresh:
            if latest.sha256 != sha256_bytes(content):
                raise SourceError(f"data/manual/{ref.file} has changed since it was first used (snapshot {latest.id}). "
                                  "Rebuild with --refresh to use the new version.")
            return latest
        return self.cache.store(key, parts[-1], content,
                                {"source": self.name, "original_file": f"data/manual/{ref.file}",
                                 "note_file": f"data/manual/{note.name}",
                                 "note_sha256": sha256_bytes(note.read_bytes())})

    def _rows(self, snap: Snapshot, ref: SourceRef) -> list[dict[str, Any]]:
        if snap.file.suffix.lower() == ".xlsx":
            book = load_workbook(io.BytesIO(snap.read_bytes()), read_only=True, data_only=True)
            sheet = book[ref.filters["sheet"]] if "sheet" in ref.filters else book.worksheets[0]
            values = list(sheet.iter_rows(values_only=True))
            book.close()
            if not values:
                return []
            header = [str(h).strip() if h is not None else "" for h in values[0]]
            return [dict(zip(header, row)) for row in values[1:]]
        return list(csv.DictReader(io.StringIO(snap.read_bytes().decode("utf-8-sig"))))

    def fetch(self, ref: SourceRef, countries: Iterable[str], refresh: bool = False) -> FetchedSeries:
        unknown = sorted(set(ref.filters) - {"sheet"})
        if unknown:
            raise SourceError(f"manual files take only the sheet filter; not {', '.join(unknown)}")
        snap = self._snapshot(ref, refresh)
        rows = self._rows(snap, ref)
        if rows and not {"iso3", "year", ref.code} <= set(rows[0]):
            raise SourceError(f"data/manual/{ref.file} needs columns iso3, year and {ref.code}; "
                              f"it has {', '.join(map(str, rows[0]))}")
        wanted = set(countries)
        records = []
        for n, row in enumerate(rows, start=2):
            iso3 = str(row["iso3"] or "").strip().upper()
            if iso3 not in wanted:
                continue
            year_text = str(row["year"] or "").strip()
            if not re.fullmatch(r"\d{4}(\.0)?", year_text):
                raise SourceError(f"data/manual/{ref.file} row {n}: year {year_text!r} is not a year")
            raw = row[ref.code]
            text = "" if raw is None else str(raw).strip()
            try:
                value = float(text) if text else math.nan
            except ValueError:
                raise SourceError(f"data/manual/{ref.file} row {n}: {ref.code} value {text!r} is not a number") from None
            records.append((iso3, int(float(year_text)), value, str(row.get("note") or "").strip()))
        unit = UnitInfo(ref.unit, ref.price_basis, ref.base_year, ref.currency)
        info = IndicatorInfo(self.name, ref.code, f"{ref.code} from data/manual/{ref.file}",
                             ref.definition or ref.note or f"See data/manual/{ref.file}.source.md", unit,
                             {"warnings": [UNCHECKED]})
        return FetchedSeries(info, make_frame(records), [snap], snap)
